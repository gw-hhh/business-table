"""Read actual exported workbooks with an independent, local openpyxl parser.

This is additional validation, not a replacement for the optional artifact_tool
rendering test. Run test:all first so browser-download fixtures are up to date.
"""
from pathlib import Path
import json
import subprocess
import zipfile
from xml.etree import ElementTree as ET
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'tests/artifacts'
OUT.mkdir(parents=True, exist_ok=True)
# Generate the two special numeric cases using the unchanged production exporter.
subprocess.run(['node', '-e', r'''
const fs=require('node:fs'),C=require('./src/core.js'),Config=require('./src/config.js'),E=require('./src/export-data.js'),W=require('./src/xlsx.js');
const s=C.normalizeSnapshot(),c=s.columns.find(c=>c.key==='amountCents');
c.number={enabled:true,mode:'percent',percentBase:'hundred',minDigits:2,maxDigits:2};
fs.writeFileSync('tests/artifacts/rules-percent.xlsx',W.write(E.createWorkbook([{...Config.SEED_ROWS[0],amountCents:1234}],s,{keys:['id','amountCents'],valueMode:'both'})));
c.number={enabled:true,mode:'decimal',scale:10000,minDigits:2,maxDigits:2,suffix:'万元'};
fs.writeFileSync('tests/artifacts/rules-scaled.xlsx',W.write(E.createWorkbook(Config.SEED_ROWS,s,{keys:['id','amountCents'],total:true})));
'''], cwd=ROOT, check=True)

specs = [
    ('browser-export.xlsx', '报价列表', 'J7'),
    ('browser-template-blank.xlsx', '数据填写', 'J5'),
    ('browser-template-example.xlsx', '示例', 'J2'),
    ('rules-mapping-export.xlsx', '报价列表', 'K7'),
    ('rules-percent.xlsx', '报价列表', 'C2'),
    ('rules-scaled.xlsx', '报价列表', 'B8'),
]
report = []
for name, sheet, last in specs:
    path = OUT / name
    if not path.is_file():
        raise FileNotFoundError(f'{name} missing. Run npm run test:all first.')
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None, f'{name}: bad ZIP CRC'
        parts = [n for n in archive.namelist() if n.endswith(('.xml', '.rels'))]
        for part in parts:
            ET.fromstring(archive.read(part))
    book = load_workbook(path, data_only=False)
    cached = load_workbook(path, data_only=True)
    ws = book[sheet]
    values = [[cell.value for cell in row] for row in cached[sheet][f'A1:{last}']]
    assert values[0][0] == '报价编号', (name, values[0])
    if name == 'browser-export.xlsx':
        assert values[1][0] == 'Q20260914-0001'
        assert abs(float(values[1][3]) - 63860) < 1e-7
        assert ws['A2'].data_type == 's' and ws['D2'].data_type == 'n'
    elif name == 'browser-template-blank.xlsx':
        assert all(v in (None, '') for row in values[1:] for v in row)
        assert len(book.worksheets) == 2
    elif name == 'browser-template-example.xlsx':
        assert len(book.worksheets) == 3 and values[1][0]
    elif name == 'rules-mapping-export.xlsx':
        assert '待处理' in values[1] and 'draft' in values[1]
        assert any('原值' in str(v) for v in values[0])
    elif name == 'rules-percent.xlsx':
        assert abs(float(values[1][1]) - .1234) < 1e-12
        assert abs(float(values[1][2]) - 12.34) < 1e-10
        assert '%' in ws['B2'].number_format and ws['B2'].data_type == 'n'
    elif name == 'rules-scaled.xlsx':
        assert abs(float(values[1][1]) - 6.386) < 1e-8
        assert abs(float(values[7][1]) - 102.645) < 1e-8
        assert ws['B8'].data_type == 'f'
    report.append({'file': name, 'zip_crc': 'passed', 'xml_parts': len(parts),
                   'independent_parser': 'openpyxl', 'import': 'passed',
                   'first_row': values[0], 'render': 'not_performed'})
    book.close()
    cached.close()
    print('PASS', name, flush=True)
(OUT / 'independent-xlsx-openpyxl-report.json').write_text(
    json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(f'RESULT {len(report)}/{len(specs)} passed; rendering not performed')
